from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Document:
    text: str
    name: str = "texte"
    # (char offset where the page starts, page number); empty when the format has no pages.
    pages: list[tuple[int, int]] = field(default_factory=list)

    def page_at(self, char: int) -> int | None:
        page = None
        for start, number in self.pages:
            if start > char:
                break
            page = number
        return page


def clean_extracted(text: str) -> str:
    """Undo PDF line wrapping: the models must see prose, not layout, or everything looks 'human'."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "")
    text = re.sub(r"(\w)-\n(?=[a-zà-ÿ])", r"\1", text)
    text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_document(path: str | Path) -> Document:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        parts, pages, pos = [], [], 0
        for number, page in enumerate(PdfReader(str(path)).pages, start=1):
            chunk = clean_extracted(page.extract_text() or "")
            if not chunk:
                continue
            pages.append((pos, number))
            parts.append(chunk)
            pos += len(chunk) + 2
        return Document("\n\n".join(parts), path.name, pages)
    if suffix == ".docx":
        from docx import Document as Docx

        paragraphs = [p.text.strip() for p in Docx(str(path)).paragraphs]
        return Document("\n\n".join(p for p in paragraphs if p), path.name)
    if suffix in {".txt", ".md", ".markdown", ".rst", ".text", ""}:
        return Document(read_text_file(path).replace("\r\n", "\n").strip(), path.name)
    raise ValueError(f"Format non pris en charge : {suffix} (pdf, docx, txt, md)")


_SENTENCE_END = re.compile(r"(?<=[.!?…»\"”])\s+(?=[A-ZÀ-ÖØ-Þ«\"“(\d—–-])|\n\s*\n")


def sentence_spans(text: str) -> list[tuple[int, int]]:
    """Character spans of sentences; paragraph breaks always end a sentence."""
    spans, start = [], 0
    for m in _SENTENCE_END.finditer(text):
        if text[start : m.start()].strip():
            spans.append((start, m.start()))
        start = m.end()
    if text[start:].strip():
        spans.append((start, len(text)))
    return spans
